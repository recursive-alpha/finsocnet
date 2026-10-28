"""Adapters that need no network: documents, YouTube parsing, Whisper (fake model), web search (fake ddgs)."""

from __future__ import annotations

import json
import types
from pathlib import Path

import pytest

from finsocnet.core import transcribe
from finsocnet.core.records import schema_errors
from tests.integration.test_adapters_api import valid

FX = Path(__file__).resolve().parents[1] / "fixtures"


def make_pdf(path: Path, pages: list[str]) -> None:
    import pymupdf
    doc = pymupdf.open()
    for t in pages:
        page = doc.new_page()
        page.insert_text((72, 72), t)
    doc.save(path)


def test_pdf_pages_become_page_segments_and_books_are_detected(repo, tmp_path):
    from finsocnet.adapters.document import build_record
    p = tmp_path / "paper.pdf"
    make_pdf(p, ["Chapter one: position sizing keeps you alive in drawdowns.", "", "Chapter two: never average down on a losing trade."])
    rec = valid(build_record(p, title="Sizing", author="A. Author", ocr="off"))
    assert rec["id"].startswith("file:") and rec["content_type"] == "file" and [s["locator"]["page"] for s in rec["segments"]] == [1, 3]
    assert rec["extra"]["empty_pages"] == 1 and rec["title"] == "Sizing"
    big = tmp_path / "book.pdf"
    make_pdf(big, [f"Page {i} has enough text to be a real page of a book." for i in range(85)])
    b = valid(build_record(big, source_id="book:0123456789abcdef0123456789abcdef", ocr="off"))
    assert b["content_type"] == "book" and b["id"] == "book:0123456789abcdef0123456789abcdef" and b["platform"] == "book" and len(b["segments"]) == 85
    arx = valid(build_record(p, source_id="arxiv:2401.01234", ocr="off"))
    assert arx["content_type"] == "paper" and arx["platform"] == "arxiv"


def test_docx_epub_html_and_text(repo, tmp_path):
    import docx
    from ebooklib import epub
    from finsocnet.adapters.document import build_record
    d = tmp_path / "notes.docx"
    doc = docx.Document()
    doc.core_properties.title = "My notes"
    doc.add_paragraph("Rule one: define the stop before the entry, always, without exception.")
    doc.add_paragraph("Rule two: never risk more than one percent of the account on a single idea.")
    doc.save(d)
    r = valid(build_record(d))
    assert r["title"] == "My notes" and "define the stop" in r["segments"][0]["text"]

    e = tmp_path / "book.epub"
    b = epub.EpubBook()
    b.set_identifier("x"); b.set_title("Trading Book"); b.add_author("Jane Writer"); b.set_language("en")
    c1 = epub.EpubHtml(title="Ch1", file_name="c1.xhtml", lang="en")
    c1.content = "<h1>Entries</h1><p>Buy the breakout only when volume confirms the move above resistance.</p><p>Wait for the close.</p>"
    b.add_item(c1)
    b.spine = ["nav", c1]
    b.add_item(epub.EpubNcx()); b.add_item(epub.EpubNav())
    epub.write_epub(str(e), b)
    er = valid(build_record(e))
    assert er["content_type"] == "book" and er["author"]["handle"] == "Jane Writer" and er["segments"][0]["locator"]["chapter"] == "Entries"

    h = tmp_path / "page.html"
    h.write_text("<html><body><h1>Title</h1><p>" + "Useful sentence about risk. " * 6 + "</p></body></html>")
    t = tmp_path / "notes.md"
    t.write_text("# Notes\n\n" + "The first rule of trading is capital preservation, always. " * 3 + "\n\nSecond paragraph that is also quite long and mentions $SPY and https://example.com/a for reference.")
    assert valid(build_record(h))["segments"]
    assert valid(build_record(t))["links"] == ["https://example.com/a"]
    from finsocnet.core.output import CliError
    bad = tmp_path / "x.xyz"
    bad.write_bytes(b"1")
    with pytest.raises(CliError) as e:
        build_record(bad)
    assert e.value.code == "UNSUPPORTED"


def test_scan_moves_processed_files_and_reports_unreadable(repo, conn, run_id, tmp_path, capsys):
    from finsocnet.adapters.document import Document
    import importlib.util
    inbox = repo / "data" / "inbox" / "documents"
    inbox.mkdir(parents=True)
    (inbox / "a.txt").write_text("A long enough paragraph about swing trading rules and discipline in markets.\n\nAnother paragraph that has plenty of words in it too, honestly.")
    (inbox / "b.xyz").write_bytes(b"\x00\x01")
    res = Document().scan_dir()
    assert [(p.name, err is None) for p, _r, err in res] == [("a.txt", True), ("b.xyz", False)]


def test_audio_transcription_is_cached_and_windowed(repo, tmp_path):
    seg = lambda s, e, t: types.SimpleNamespace(start=s, end=e, text=t)
    calls = []

    class FakeModel:
        def transcribe(self, path, **kw):
            calls.append(path)
            return iter([seg(0, 20, " Hello traders."), seg(20, 41, " Today is about risk."), seg(41, 90, " And exits.")]), types.SimpleNamespace(language="en", duration=90.0)

    transcribe.set_model_factory(FakeModel)
    try:
        wav = tmp_path / "talk.wav"
        wav.write_bytes(b"RIFF....fake audio bytes")
        r1 = transcribe.transcribe(str(wav))
        r2 = transcribe.transcribe(str(wav))
        assert len(calls) == 1 and r1["cached"] is False and r2["cached"] is True and r1["segments"][1]["text"] == "Today is about risk."
        from finsocnet.core.segments import transcript_windows
        w = transcript_windows(transcribe.as_entries(r1))
        assert [x["locator"]["end_sec"] for x in w] == [41.0, 90.0]  # 45 s windows
        from finsocnet.adapters.document import build_record
        rec = valid(build_record(wav))
        assert rec["extra"]["recording"] is True and rec["segments"][0]["kind"] == "transcript"
    finally:
        transcribe.set_model_factory(None)


def test_youtube_parsing_and_caption_choice():
    from finsocnet.adapters import youtube as yt
    assert yt.video_id("https://youtu.be/B1ALxAVtFfI?t=5") == "B1ALxAVtFfI" and yt.video_id("https://www.youtube.com/shorts/B1ALxAVtFfI") == "B1ALxAVtFfI"
    assert yt.choose_caption([("en", True), ("en", False), ("vi", False)], "vi") == 2
    assert yt.choose_caption([("en", True), ("vi", True)], "vi") == 1
    assert yt.choose_caption([("en", True)], "de") == 0 and yt.choose_caption([("fr", False)], "de") is None
    data = json.loads((FX / "transcripts" / "abc123def45.json").read_text())
    info = {"id": "abc123def45", "title": "Selling into strength", "description": "Rules at https://example.com/rules with $SPY and @example", "uploader_id": "@Example", "channel": "Example",
            "channel_id": "UC" + "a" * 22, "channel_url": "https://youtube.com/channel/x", "upload_date": "20260902", "duration": 600, "view_count": 10, "chapters": [{"title": "Intro", "start_time": 0}], "was_live": False}
    rec = valid(yt.video_record(info, data["transcript"], "manual", "en"))
    assert rec["id"] == "youtube:abc123def45" and rec["content_type"] == "video" and rec["published_at"] == "2026-09-02" and rec["extra"]["caption_source"] == "manual"
    assert rec["source"]["id"] == "youtube:channel:UC" + "a" * 22 and "$SPY" in rec["mentions"] and rec["links"] == ["https://example.com/rules"]
    assert all(s["locator"]["end_sec"] - s["locator"]["start_sec"] <= 52 for s in rec["segments"])
    live = yt.video_record({**info, "was_live": True, "duration": 4 * 3600}, None, None, None)
    assert live["content_type"] == "stream" and live["extra"]["long_form"] is True and live["segments"] == []


def test_web_search_maps_results(repo, monkeypatch):
    import sys
    from finsocnet.adapters.web_search import WebSearch
    class FakeDDGS:
        def text(self, q, region=None, timelimit=None, max_results=5):
            return [{"title": "Qullamaggie setups", "href": "https://qullamaggie.com/?utm_source=x", "body": "Breakouts and EPs"}]
        def news(self, *a, **k):
            return []
    monkeypatch.setitem(sys.modules, "ddgs", types.SimpleNamespace(DDGS=FakeDDGS))
    items = WebSearch().search("qullamaggie", limit=5, site="qullamaggie.com")
    from finsocnet.core import ids
    assert items[0]["id"] == ids.web_content_id("https://qullamaggie.com/") and items[0]["title"] == "Qullamaggie setups"
